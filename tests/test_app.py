"""Focused tests for connected Flask workflows and safe local persistence."""

from __future__ import annotations

import json
import tempfile
import unittest
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import app as application
from docx import Document
from utils import json_handler


class CareerPlatformTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.data_patch = patch.object(json_handler, "_DATA_DIR", Path(self.temporary.name))
        self.data_patch.start()
        self.ai_patch = patch.object(application.gemini_service, "generate_json", side_effect=application.gemini_service.GeminiError("Demo mode for test"))
        self.ai_patch.start()
        self.client = application.app.test_client()

    def tearDown(self):
        self.ai_patch.stop()
        self.data_patch.stop()
        self.temporary.cleanup()

    def create_plan(self):
        response = self.client.post("/api/career/generate", json={
            "name": "Sam Student", "target_role": "Python Developer", "current_skills": "Python, Git",
            "education": "Computer science student", "skill_level": "Beginner", "available_time": "30 minutes/day",
            "preferred_domain": "Education technology", "experience": "Volunteer tutoring", "projects": "A small calculator",
        })
        self.assertEqual(response.status_code, 200)
        return response.get_json()

    def test_health_home_and_ai_status(self):
        self.assertEqual(self.client.get("/api/health").get_json(), {"status": "ok"})
        self.assertEqual(self.client.get("/").status_code, 200)
        status = self.client.get("/api/ai-status").get_json()
        self.assertEqual(status["service"], "Gemini")
        self.assertNotIn("key", json.dumps(status).lower())

    def test_career_plan_adapts_and_stores_progress(self):
        plan = self.create_plan()
        self.assertEqual(len(plan["roadmap"]), 30)
        self.assertTrue(plan["demo"])
        self.assertTrue(all(day["estimated_minutes"] <= 30 for day in plan["roadmap"]))
        self.assertIn("Flask", [item["skill"] for item in plan["gaps"]["priorities"]])
        done = self.client.post("/api/roadmap/progress", json={"day": 1}).get_json()
        self.assertEqual(done["progress"]["completed_days"], [1])
        duplicate = self.client.post("/api/roadmap/progress", json={"day": 1}).get_json()
        self.assertEqual(duplicate["progress"]["completed_days"], [1])
        plan["created_at"] = (date.today() - timedelta(days=3)).isoformat()
        json_handler.update_json("career_plans", "local", plan)
        adapted = self.client.post("/api/roadmap/adapt", json={}).get_json()
        self.assertEqual(adapted["missed_count"], 3)
        self.assertEqual(adapted["remaining_count"], 29)
        self.assertEqual(self.client.get("/api/dashboard").get_json()["progress"]["completed_days"], [1])

    def test_required_fields_and_skill_gap_validation(self):
        missing = self.client.post("/api/career/generate", json={"name": ""})
        self.assertEqual(missing.status_code, 400)
        self.assertIn("Name is required", missing.get_json()["error"])
        gap = self.client.post("/api/skill-gap", json={"target_role": "Data Analyst", "current_skills": "SQL, Excel"}).get_json()
        self.assertIn("SQL", gap["already_have"])
        self.assertIn("Python", gap["need_to_learn"])
        self.assertEqual(self.client.get("/api/career-path?role=Python%20Developer").status_code, 200)

    def test_resume_txt_analysis_is_in_memory_and_role_specific(self):
        response = self.client.post("/api/resume/analyze", data={
            "target_role": "Python Developer",
            "resume": (BytesIO(b"alex@example.com\nSkills\nPython, SQL, Git\nEducation\nComputer Science\nProjects\nBuilt a calculator"), "resume.txt"),
        }, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 200)
        result = response.get_json()
        self.assertIn("Python", result["matching_skills"])
        self.assertIn("ATS system", result["disclaimer"])
        history = json_handler.read_json("resume_history")
        self.assertEqual(len(history), 1)
        self.assertNotIn("alex@example.com", json.dumps(history))
        invalid = self.client.post("/api/resume/analyze", data={"target_role": "Developer", "resume": (BytesIO(b"x"), "resume.exe")}, content_type="multipart/form-data")
        self.assertEqual(invalid.status_code, 400)

    def test_docx_table_detection_and_resume_issues(self):
        document = Document()
        document.add_paragraph("alex@example.com")
        document.add_heading("Skills", level=1)
        document.add_paragraph("Python, SQL")
        document.add_heading("Projects", level=1)
        document.add_paragraph("Made a small project.")
        document.add_heading("Education", level=1)
        document.add_paragraph("Computer science")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Flask"
        buffer = BytesIO()
        document.save(buffer)
        response = self.client.post("/api/resume/analyze", data={
            "target_role": "Python Developer", "resume": (BytesIO(buffer.getvalue()), "resume.docx"),
        }, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 200)
        result = response.get_json()
        problems = [issue["problem"] for issue in result["issues"]]
        self.assertIn("Table-based layout detected", problems)
        self.assertTrue(any("vague" in problem for problem in problems))
        self.assertIn("Flask", result["matching_skills"])

    def test_job_analysis_projects_and_resources(self):
        comparison = self.client.post("/api/job-description/analyze", json={
            "target_role": "Python Developer", "job_description": "Python Flask SQL REST APIs", "resume_text": "Python SQL",
        }).get_json()
        self.assertIn("Flask", comparison["missing"])
        self.assertTrue(any("genuinely have" in item for item in comparison["actions"]))
        projects = self.client.post("/api/projects/recommend", json={"target_role": "Python Developer", "current_skills": ["Python"]}).get_json()
        self.assertEqual(len(projects["projects"]), 3)
        saved = self.client.post("/api/resources/save", json={"url": "https://docs.python.org/3/tutorial/"})
        self.assertEqual(saved.status_code, 200)
        rejected = self.client.post("/api/resources/save", json={"url": "https://example.invalid/fake"})
        self.assertEqual(rejected.status_code, 400)

    def test_project_log_updates_readiness_and_history(self):
        result = self.client.post("/api/projects/log", json={
            "name": "Study planner", "technologies": "Python, Flask", "description": "Tracks study sessions.",
        })
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.get_json()["achievements"]["xp"], 40)
        self.assertGreater(self.client.get("/api/dashboard").get_json()["readiness"]["Projects"], 0)
        history = self.client.get("/api/history").get_json()
        self.assertEqual(history["projects"][0]["name"], "Study planner")

    def test_interview_answers_persist_and_completion_is_idempotent(self):
        session = self.client.post("/api/interview/start", json={"target_role": "Data Analyst"}).get_json()
        self.assertEqual(len(session["questions"]), 5)
        answer = self.client.post("/api/interview/answer", json={
            "session_id": session["id"], "question_index": 0,
            "question": session["questions"][0]["question"], "answer": "I used SQL because it helped. For example, I tested the query.",
        })
        self.assertEqual(answer.status_code, 200)
        saved = json_handler.read_json("interview_history")[0]
        self.assertEqual(len(saved["answers"]), 1)
        first = self.client.post("/api/interview/complete", json={"session_id": session["id"]}).get_json()
        second = self.client.post("/api/interview/complete", json={"session_id": session["id"]}).get_json()
        self.assertEqual(first["achievements"]["xp"], second["achievements"]["xp"])

    def test_daily_challenge_cannot_be_claimed_twice(self):
        self.create_plan()
        started = self.client.post("/api/challenge/start", json={})
        self.assertEqual(started.status_code, 200)
        first = self.client.post("/api/challenge/complete", json={}).get_json()
        second = self.client.post("/api/challenge/complete", json={}).get_json()
        self.assertEqual(first["achievements"]["xp"], 20)
        self.assertEqual(second["achievements"]["xp"], 20)


if __name__ == "__main__":
    unittest.main()
