"""In-memory extraction for supported resume formats."""

from __future__ import annotations

from io import BytesIO
import re


def extract_text(filename: str, content: bytes) -> str:
    extension = filename.rsplit(".", 1)[-1].lower()
    if extension == "txt":
        return content.decode("utf-8", errors="replace")
    if extension == "pdf":
        from PyPDF2 import PdfReader

        reader = PdfReader(BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if extension == "docx":
        from docx import Document

        document = Document(BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        tables = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        return "\n".join(paragraphs + tables)
    raise ValueError("Only PDF, DOCX, and TXT files are supported.")


def inspect_layout(filename: str, content: bytes) -> dict[str, int | bool]:
    extension = filename.rsplit(".", 1)[-1].lower()
    if extension == "docx":
        from docx import Document

        document = Document(BytesIO(content))
        font_sizes = {run.font.size.pt for paragraph in document.paragraphs for run in paragraph.runs if run.font.size}
        font_names = {run.font.name for paragraph in document.paragraphs for run in paragraph.runs if run.font.name}
        return {"tables": len(document.tables), "images": len(document.inline_shapes), "excessive_formatting": len(font_sizes) > 5 or len(font_names) > 4}
    if extension == "pdf":
        from PyPDF2 import PdfReader

        reader = PdfReader(BytesIO(content))
        has_images = False
        for page in reader.pages:
            resources = page.get("/Resources")
            if resources and resources.get_object().get("/XObject"):
                has_images = True
                break
        return {"tables": 0, "images": int(has_images), "excessive_formatting": False}
    return {"tables": 0, "images": 0, "excessive_formatting": False}


def inspect_text(text: str) -> list[dict[str, str]]:
    issues = []

    def add(problem: str, why: str, improve: str) -> None:
        issues.append({"problem": problem, "why": why, "improve": improve})

    if not re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", text) and not re.search(r"(?:\+?\d[\d ().-]{7,}\d)", text):
        add("Contact details were not detected", "A recruiter needs a way to reach you.", "Add a professional email address or phone number if appropriate.")
    headings = ["experience", "education", "skills", "projects"]
    if sum(bool(re.search(rf"\b{heading}\b", text, re.I)) for heading in headings) < 3:
        add("Few standard section headings were detected", "Clear structure makes relevant information easier to scan.", "Use straightforward headings such as Skills, Projects, Education, and Experience.")
    if any(len(paragraph) > 600 for paragraph in text.splitlines()):
        add("A long paragraph was detected", "Dense blocks of text can be difficult to scan.", "Break the paragraph into concise bullets, keeping the original facts.")
    if re.search(r"\b(?:worked on|responsible for|helped with|made)\b", text, re.I):
        add("Some project or experience wording is vague", "Generic descriptions do not show your specific contribution.", "Explain what you personally did, using details you can support.")
    if not re.search(r"\b\d+(?:\.\d+)?\s*%\b|\b\d+\s+(?:users|records|hours|weeks|days|requests)\b", text, re.I):
        add("No measurable outcomes were detected", "Specific scale can help readers understand impact when it is available.", "Add a verifiable measure only when you can support it; do not estimate or invent numbers.")
    date_styles = [
        bool(re.search(r"\b(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)\s+\d{4}\b", text, re.I)),
        bool(re.search(r"\b\d{1,2}[/-]\d{4}\b", text)),
        bool(re.search(r"\b(?:19|20)\d{2}[/-]\d{1,2}\b", text)),
    ]
    if sum(date_styles) > 1:
        add("More than one date format was detected", "Inconsistent date styles can make a resume harder to follow.", "Choose one date format and use it consistently.")
    return issues
